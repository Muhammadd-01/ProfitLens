"""File processing utilities for data engineering and ingestion.

Handles:
- Encoding detection (UTF-8, Latin-1, Windows-1252) via chardet
- Delimiter sniffing (comma, semicolon, tab, pipe) via csv.Sniffer
- Safe chunked streaming to avoid memory exhaustion (OOM)
- Fast header & row count extraction for CSV and XLSX files
"""

from __future__ import annotations

import csv
import io
import os
import chardet
from typing import Tuple, List, Dict, Any, Optional
import openpyxl


def detect_encoding(sample_bytes: bytes) -> str:
    """Detect text encoding from raw bytes.
    
    Checks for common BOM headers first, then falls back to chardet.
    Defaults to utf-8 if confidence is low.
    """
    if sample_bytes.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig"
    if sample_bytes.startswith(b"\xff\xfe"):
        return "utf-16-le"
    if sample_bytes.startswith(b"\xfe\xff"):
        return "utf-16-be"

    result = chardet.detect(sample_bytes)
    detected = result.get("encoding")
    confidence = result.get("confidence", 0.0)

    if detected and confidence >= 0.7:
        det = detected.lower()
        if det == "ascii":
            return "utf-8"
        return det
    return "utf-8"


def detect_delimiter(sample_text: str) -> str:
    """Detect the CSV delimiter (comma, tab, semicolon, pipe) using csv.Sniffer."""
    try:
        dialect = csv.Sniffer().sniff(sample_text, delimiters=[",", ";", "\t", "|"])
        return dialect.delimiter
    except Exception:
        # Default fallback is comma
        return ","


def inspect_csv_file(file_path: str) -> Tuple[int, List[str], str, str]:
    """Inspect a CSV file without loading the full file into memory.
    
    Returns:
        (row_count, column_names, encoding, delimiter)
    """
    # 1. Read first 64KB to detect encoding and delimiter
    with open(file_path, "rb") as f:
        sample_bytes = f.read(65536)

    encoding = detect_encoding(sample_bytes)

    # Decode sample text
    try:
        sample_text = sample_bytes.decode(encoding)
    except Exception:
        encoding = "latin-1"
        sample_text = sample_bytes.decode("latin-1", errors="replace")

    delimiter = detect_delimiter(sample_text)

    # 2. Extract header and count rows via streaming
    columns: List[str] = []
    row_count = 0

    with open(file_path, "r", encoding=encoding, errors="replace") as f:
        reader = csv.reader(f, delimiter=delimiter)
        try:
            raw_header = next(reader)
            # Clean header column names: strip whitespace, remove empty quotes
            columns = [col.strip() for col in raw_header if col.strip()]
        except StopIteration:
            return 0, [], encoding, delimiter

        # Count remaining rows
        for _ in reader:
            row_count += 1

    return row_count, columns, encoding, delimiter


def inspect_excel_file(file_path: str) -> Tuple[int, List[str]]:
    """Inspect an Excel (.xlsx) file in read-only streaming mode.
    
    Returns:
        (row_count, column_names)
    """
    workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    first_sheet_name = workbook.sheetnames[0]
    sheet = workbook[first_sheet_name]

    columns: List[str] = []
    row_count = 0

    for i, row in enumerate(sheet.iter_rows(values_only=True)):
        if i == 0:
            columns = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
        else:
            # Only count non-empty rows
            if any(cell is not None for cell in row):
                row_count += 1

    workbook.close()
    return row_count, columns
